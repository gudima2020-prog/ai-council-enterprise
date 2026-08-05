from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from enum import Enum
import hashlib
from io import BytesIO
import json
import math
from typing import Any, Iterable
import unicodedata

from docx import Document as DocxDocument
from docx.document import Document as _DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pypdf import PdfReader

from backend.documents.intake import (
    DocumentFormat,
    DocumentIntakeDescriptor,
)


class DocumentExtractionError(RuntimeError):
    def __init__(
        self,
        code: str,
        message: str,
        *,
        details: dict[str, Any] | None = None,
    ) -> None:
        super().__init__(message)
        self.code = code
        self.details = dict(details or {})


class ExtractionUnitKind(str, Enum):
    PDF_PAGE = "pdf_page"
    DOCX_PARAGRAPH = "docx_paragraph"
    DOCX_TABLE_CELL = "docx_table_cell"
    XLSX_CELL = "xlsx_cell"
    TXT_DOCUMENT = "txt_document"


@dataclass(frozen=True, slots=True)
class DocumentExtractionPolicy:
    max_total_characters: int = 5_000_000
    max_unit_characters: int = 250_000
    max_units: int = 100_000
    max_pdf_pages: int = 5_000
    max_docx_paragraphs: int = 100_000
    max_docx_tables: int = 10_000
    max_docx_cells: int = 500_000
    max_xlsx_sheets: int = 256
    max_xlsx_rows_per_sheet: int = 200_000
    max_xlsx_cells: int = 1_000_000

    def __post_init__(self) -> None:
        for name in self.__dataclass_fields__:
            if getattr(self, name) <= 0:
                raise ValueError(f"{name} must be positive")


@dataclass(frozen=True, slots=True)
class ExtractedTextUnit:
    ordinal: int
    kind: ExtractionUnitKind
    text: str
    text_sha256: str
    provenance: dict[str, Any] = field(default_factory=dict)

    def to_public_dict(self) -> dict[str, Any]:
        return {
            "ordinal": self.ordinal,
            "kind": self.kind.value,
            "text": self.text,
            "text_sha256": self.text_sha256,
            "provenance": dict(self.provenance),
        }


@dataclass(frozen=True, slots=True)
class DocumentExtractionResult:
    document_format: DocumentFormat
    parser: str
    parser_version: str
    source_sha256: str
    extracted_text_sha256: str
    total_characters: int
    unit_count: int
    units: tuple[ExtractedTextUnit, ...]
    warnings: tuple[str, ...] = field(default_factory=tuple)

    def to_public_dict(
        self,
        *,
        include_text: bool = True,
    ) -> dict[str, Any]:
        units: list[dict[str, Any]] = []
        for unit in self.units:
            item = unit.to_public_dict()
            if not include_text:
                item.pop("text", None)
            units.append(item)
        return {
            "document_format": self.document_format.value,
            "parser": self.parser,
            "parser_version": self.parser_version,
            "source_sha256": self.source_sha256,
            "extracted_text_sha256": self.extracted_text_sha256,
            "total_characters": self.total_characters,
            "unit_count": self.unit_count,
            "warnings": list(self.warnings),
            "units": units,
        }


class DeterministicDocumentExtractor:
    PARSER_VERSION = "p3-001.3a-v1"

    def __init__(
        self,
        policy: DocumentExtractionPolicy | None = None,
    ) -> None:
        self._policy = policy or DocumentExtractionPolicy()

    def extract(
        self,
        *,
        descriptor: DocumentIntakeDescriptor,
        content: bytes,
    ) -> DocumentExtractionResult:
        source_sha256 = hashlib.sha256(content).hexdigest()
        if source_sha256 != descriptor.content_sha256:
            raise DocumentExtractionError(
                "DOCUMENT_EXTRACTION_SOURCE_HASH_MISMATCH",
                "Document bytes do not match the intake descriptor.",
            )

        if descriptor.document_format is DocumentFormat.PDF:
            parser = "pypdf"
            units, warnings = self._extract_pdf(content)
        elif descriptor.document_format is DocumentFormat.DOCX:
            parser = "python-docx"
            units, warnings = self._extract_docx(content)
        elif descriptor.document_format is DocumentFormat.XLSX:
            parser = "openpyxl"
            units, warnings = self._extract_xlsx(content)
        elif descriptor.document_format is DocumentFormat.TXT:
            parser = "builtin-text"
            units, warnings = self._extract_txt(
                content,
                descriptor.content_encoding,
            )
        else:
            raise DocumentExtractionError(
                "DOCUMENT_EXTRACTION_FORMAT_UNSUPPORTED",
                "Document format is not supported for extraction.",
            )

        self._enforce_result_limits(units)
        canonical_text = self._canonical_document_text(units)
        extracted_hash = hashlib.sha256(
            canonical_text.encode("utf-8")
        ).hexdigest()

        return DocumentExtractionResult(
            document_format=descriptor.document_format,
            parser=parser,
            parser_version=self.PARSER_VERSION,
            source_sha256=source_sha256,
            extracted_text_sha256=extracted_hash,
            total_characters=sum(len(unit.text) for unit in units),
            unit_count=len(units),
            units=tuple(units),
            warnings=tuple(warnings),
        )

    def _extract_pdf(
        self,
        content: bytes,
    ) -> tuple[list[ExtractedTextUnit], list[str]]:
        try:
            reader = PdfReader(BytesIO(content), strict=True)
        except Exception as exc:
            raise DocumentExtractionError(
                "DOCUMENT_EXTRACTION_PDF_INVALID",
                "PDF cannot be parsed safely.",
            ) from exc

        if reader.is_encrypted:
            raise DocumentExtractionError(
                "DOCUMENT_EXTRACTION_PDF_ENCRYPTED",
                "Encrypted PDF extraction is not supported.",
            )

        page_count = len(reader.pages)
        if page_count > self._policy.max_pdf_pages:
            raise DocumentExtractionError(
                "DOCUMENT_EXTRACTION_PDF_PAGE_LIMIT",
                "PDF page count exceeds the configured limit.",
                details={"page_count": page_count},
            )

        units: list[ExtractedTextUnit] = []
        warnings: list[str] = []
        for page_index, page in enumerate(reader.pages):
            try:
                raw_text = page.extract_text() or ""
            except Exception as exc:
                raise DocumentExtractionError(
                    "DOCUMENT_EXTRACTION_PDF_PAGE_FAILED",
                    "PDF page text extraction failed.",
                    details={"page_number": page_index + 1},
                ) from exc

            text = self._normalize_text(raw_text)
            if not text:
                warnings.append(
                    f"PDF_PAGE_{page_index + 1}_NO_TEXT"
                )
            units.append(
                self._make_unit(
                    kind=ExtractionUnitKind.PDF_PAGE,
                    text=text,
                    provenance={
                        "page_number": page_index + 1,
                    },
                    ordinal=len(units),
                )
            )
        return units, warnings

    def _extract_docx(
        self,
        content: bytes,
    ) -> tuple[list[ExtractedTextUnit], list[str]]:
        try:
            document = DocxDocument(BytesIO(content))
        except Exception as exc:
            raise DocumentExtractionError(
                "DOCUMENT_EXTRACTION_DOCX_INVALID",
                "DOCX cannot be parsed safely.",
            ) from exc

        units: list[ExtractedTextUnit] = []
        paragraph_index = 0
        table_index = 0
        cell_count = 0

        for block in self._iter_docx_blocks(document):
            if isinstance(block, Paragraph):
                paragraph_index += 1
                if (
                    paragraph_index
                    > self._policy.max_docx_paragraphs
                ):
                    raise DocumentExtractionError(
                        "DOCUMENT_EXTRACTION_DOCX_PARAGRAPH_LIMIT",
                        "DOCX paragraph count exceeds the configured limit.",
                    )
                text = self._normalize_text(block.text)
                if text:
                    units.append(
                        self._make_unit(
                            kind=(
                                ExtractionUnitKind.DOCX_PARAGRAPH
                            ),
                            text=text,
                            provenance={
                                "paragraph_index": paragraph_index,
                                "style": (
                                    block.style.name
                                    if block.style is not None
                                    else None
                                ),
                            },
                            ordinal=len(units),
                        )
                    )
                continue

            table_index += 1
            if table_index > self._policy.max_docx_tables:
                raise DocumentExtractionError(
                    "DOCUMENT_EXTRACTION_DOCX_TABLE_LIMIT",
                    "DOCX table count exceeds the configured limit.",
                )

            for row_index, row in enumerate(
                block.rows,
                start=1,
            ):
                for column_index, cell in enumerate(
                    row.cells,
                    start=1,
                ):
                    cell_count += 1
                    if (
                        cell_count
                        > self._policy.max_docx_cells
                    ):
                        raise DocumentExtractionError(
                            "DOCUMENT_EXTRACTION_DOCX_CELL_LIMIT",
                            "DOCX cell count exceeds the configured limit.",
                        )
                    text = self._normalize_text(cell.text)
                    if not text:
                        continue
                    units.append(
                        self._make_unit(
                            kind=(
                                ExtractionUnitKind.DOCX_TABLE_CELL
                            ),
                            text=text,
                            provenance={
                                "table_index": table_index,
                                "row_index": row_index,
                                "column_index": column_index,
                            },
                            ordinal=len(units),
                        )
                    )

        warnings = (
            ["DOCX_NO_EXTRACTABLE_TEXT"]
            if not units
            else []
        )
        return units, warnings

    def _extract_xlsx(
        self,
        content: bytes,
    ) -> tuple[list[ExtractedTextUnit], list[str]]:
        try:
            workbook = load_workbook(
                filename=BytesIO(content),
                read_only=True,
                data_only=False,
                keep_links=False,
            )
        except Exception as exc:
            raise DocumentExtractionError(
                "DOCUMENT_EXTRACTION_XLSX_INVALID",
                "XLSX cannot be parsed safely.",
            ) from exc

        try:
            if (
                len(workbook.worksheets)
                > self._policy.max_xlsx_sheets
            ):
                raise DocumentExtractionError(
                    "DOCUMENT_EXTRACTION_XLSX_SHEET_LIMIT",
                    "XLSX sheet count exceeds the configured limit.",
                )

            units: list[ExtractedTextUnit] = []
            cell_count = 0
            for sheet_index, sheet in enumerate(
                workbook.worksheets,
                start=1,
            ):
                for row_index, row in enumerate(
                    sheet.iter_rows(),
                    start=1,
                ):
                    if (
                        row_index
                        > self._policy.max_xlsx_rows_per_sheet
                    ):
                        raise DocumentExtractionError(
                            "DOCUMENT_EXTRACTION_XLSX_ROW_LIMIT",
                            "XLSX row count exceeds the configured limit.",
                            details={
                                "sheet_name": sheet.title,
                            },
                        )
                    for cell in row:
                        if cell.value is None:
                            continue
                        cell_count += 1
                        if (
                            cell_count
                            > self._policy.max_xlsx_cells
                        ):
                            raise DocumentExtractionError(
                                "DOCUMENT_EXTRACTION_XLSX_CELL_LIMIT",
                                "XLSX cell count exceeds the configured limit.",
                            )
                        text = self._normalize_text(
                            self._stable_cell_value(
                                cell.value
                            )
                        )
                        if not text:
                            continue
                        units.append(
                            self._make_unit(
                                kind=ExtractionUnitKind.XLSX_CELL,
                                text=text,
                                provenance={
                                    "sheet_index": sheet_index,
                                    "sheet_name": sheet.title,
                                    "row_index": cell.row,
                                    "column_index": cell.column,
                                    "coordinate": cell.coordinate,
                                    "data_type": cell.data_type,
                                },
                                ordinal=len(units),
                            )
                        )
        finally:
            workbook.close()

        warnings = (
            ["XLSX_NO_EXTRACTABLE_VALUES"]
            if not units
            else []
        )
        return units, warnings

    def _extract_txt(
        self,
        content: bytes,
        content_encoding: str | None,
    ) -> tuple[list[ExtractedTextUnit], list[str]]:
        encoding = (content_encoding or "utf-8").casefold()
        aliases = {
            "utf-8-sig": "utf-8-sig",
            "utf-8": "utf-8",
            "utf-16": "utf-16",
            "utf-16-le": "utf-16-le",
            "utf-16-be": "utf-16-be",
            "windows-1251": "cp1251",
            "cp1251": "cp1251",
        }
        codec = aliases.get(encoding)
        if codec is None:
            raise DocumentExtractionError(
                "DOCUMENT_EXTRACTION_TEXT_ENCODING_UNSUPPORTED",
                "TXT encoding is not supported.",
                details={"encoding": content_encoding},
            )

        try:
            text = content.decode(codec, errors="strict")
        except UnicodeDecodeError as exc:
            raise DocumentExtractionError(
                "DOCUMENT_EXTRACTION_TEXT_DECODE_FAILED",
                "TXT content cannot be decoded with the intake encoding.",
            ) from exc

        normalized = self._normalize_text(text)
        units = [
            self._make_unit(
                kind=ExtractionUnitKind.TXT_DOCUMENT,
                text=normalized,
                provenance={"encoding": encoding},
                ordinal=0,
            )
        ]
        warnings = (
            ["TXT_EMPTY_AFTER_NORMALIZATION"]
            if not normalized
            else []
        )
        return units, warnings

    def _make_unit(
        self,
        *,
        kind: ExtractionUnitKind,
        text: str,
        provenance: dict[str, Any],
        ordinal: int,
    ) -> ExtractedTextUnit:
        if len(text) > self._policy.max_unit_characters:
            raise DocumentExtractionError(
                "DOCUMENT_EXTRACTION_UNIT_TOO_LARGE",
                "Extracted text unit exceeds the configured limit.",
                details={
                    "kind": kind.value,
                    "characters": len(text),
                },
            )
        return ExtractedTextUnit(
            ordinal=ordinal,
            kind=kind,
            text=text,
            text_sha256=hashlib.sha256(
                text.encode("utf-8")
            ).hexdigest(),
            provenance=self._canonical_provenance(
                provenance
            ),
        )

    def _enforce_result_limits(
        self,
        units: list[ExtractedTextUnit],
    ) -> None:
        if len(units) > self._policy.max_units:
            raise DocumentExtractionError(
                "DOCUMENT_EXTRACTION_UNIT_LIMIT",
                "Extracted unit count exceeds the configured limit.",
            )

        total = sum(len(unit.text) for unit in units)
        if total > self._policy.max_total_characters:
            raise DocumentExtractionError(
                "DOCUMENT_EXTRACTION_CHARACTER_LIMIT",
                "Extracted text exceeds the configured limit.",
                details={"characters": total},
            )

    @staticmethod
    def _normalize_text(value: str) -> str:
        normalized = unicodedata.normalize(
            "NFKC",
            value,
        )
        normalized = normalized.replace(
            "\r\n",
            "\n",
        ).replace("\r", "\n")
        lines = [
            line.rstrip()
            for line in normalized.split("\n")
        ]
        while lines and not lines[0]:
            lines.pop(0)
        while lines and not lines[-1]:
            lines.pop()
        return "\n".join(lines)

    @staticmethod
    def _stable_cell_value(value: Any) -> str:
        if isinstance(value, bool):
            return "TRUE" if value else "FALSE"
        if isinstance(value, (datetime, date, time)):
            return value.isoformat()
        if isinstance(value, Decimal):
            return format(value, "f")
        if isinstance(value, float):
            if not math.isfinite(value):
                raise DocumentExtractionError(
                    "DOCUMENT_EXTRACTION_XLSX_NUMBER_INVALID",
                    "XLSX contains a non-finite numeric value.",
                )
            return format(value, ".15g")
        if isinstance(value, int):
            return str(value)
        if isinstance(value, str):
            return value
        return str(value)

    @staticmethod
    def _canonical_provenance(
        value: dict[str, Any],
    ) -> dict[str, Any]:
        serialized = json.dumps(
            value,
            ensure_ascii=True,
            sort_keys=True,
            separators=(",", ":"),
        )
        result = json.loads(serialized)
        if not isinstance(result, dict):
            raise ValueError(
                "Provenance must serialize to an object."
            )
        return result

    @staticmethod
    def _canonical_document_text(
        units: Iterable[ExtractedTextUnit],
    ) -> str:
        return "\n\n".join(
            unit.text
            for unit in units
        )

    @staticmethod
    def _iter_docx_blocks(
        parent: _DocxDocument,
    ) -> Iterable[Paragraph | Table]:
        body = parent.element.body
        for child in body.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, parent)
            elif isinstance(child, CT_Tbl):
                yield Table(child, parent)
