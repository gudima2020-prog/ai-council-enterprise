from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook
from pypdf import PdfWriter

from backend.documents.extraction import (
    DeterministicDocumentExtractor,
    DocumentExtractionError,
    DocumentExtractionPolicy,
    ExtractionUnitKind,
)
from backend.documents.intake import (
    DocumentIntakeRequest,
    DocumentIntakeService,
)
from backend.runtime_policy import DataClassification


def intake(
    *,
    filename: str,
    content_type: str,
    content: bytes,
):
    return DocumentIntakeService().inspect_bytes(
        DocumentIntakeRequest(
            workspace_id="workspace_alpha",
            filename=filename,
            content_type=content_type,
            classification=DataClassification.CONFIDENTIAL,
        ),
        content,
    )


def make_docx() -> bytes:
    document = DocxDocument()
    document.add_paragraph("Heading")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "B1"
    table.cell(1, 0).text = "A2"
    table.cell(1, 1).text = "B2"
    document.add_paragraph("Tail")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = "Name"
    sheet["B1"] = "Value"
    sheet["A2"] = "Alpha"
    sheet["B2"] = 12.5
    formula = workbook.create_sheet("Formula")
    formula["A1"] = "=1+1"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def test_txt_extraction_is_normalized_and_hashed() -> None:
    content = b"  first\r\nsecond  \r\n"
    descriptor = intake(
        filename="notes.txt",
        content_type="text/plain",
        content=content,
    )

    result = DeterministicDocumentExtractor().extract(
        descriptor=descriptor,
        content=content,
    )

    assert result.parser == "builtin-text"
    assert result.unit_count == 1
    assert result.units[0].kind == (
        ExtractionUnitKind.TXT_DOCUMENT
    )
    assert result.units[0].text == "  first\nsecond"
    assert result.total_characters == len(
        "  first\nsecond"
    )
    assert len(result.extracted_text_sha256) == 64


def test_txt_cp1251_uses_intake_encoding() -> None:
    content = "Привет".encode("cp1251")
    descriptor = intake(
        filename="notes.txt",
        content_type="text/plain",
        content=content,
    )

    result = DeterministicDocumentExtractor().extract(
        descriptor=descriptor,
        content=content,
    )

    assert result.units[0].text == "Привет"
    assert result.units[0].provenance[
        "encoding"
    ] in {"windows-1251", "cp1251"}


def test_source_hash_mismatch_is_rejected() -> None:
    content = b"hello"
    descriptor = intake(
        filename="notes.txt",
        content_type="text/plain",
        content=content,
    )

    with pytest.raises(
        DocumentExtractionError,
    ) as captured:
        DeterministicDocumentExtractor().extract(
            descriptor=descriptor,
            content=b"other",
        )

    assert captured.value.code == (
        "DOCUMENT_EXTRACTION_SOURCE_HASH_MISMATCH"
    )


def test_docx_preserves_paragraph_and_table_provenance() -> None:
    content = make_docx()
    descriptor = intake(
        filename="report.docx",
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        content=content,
    )

    result = DeterministicDocumentExtractor().extract(
        descriptor=descriptor,
        content=content,
    )

    assert [
        unit.kind
        for unit in result.units
    ] == [
        ExtractionUnitKind.DOCX_PARAGRAPH,
        ExtractionUnitKind.DOCX_TABLE_CELL,
        ExtractionUnitKind.DOCX_TABLE_CELL,
        ExtractionUnitKind.DOCX_TABLE_CELL,
        ExtractionUnitKind.DOCX_TABLE_CELL,
        ExtractionUnitKind.DOCX_PARAGRAPH,
    ]
    assert result.units[0].text == "Heading"
    assert result.units[1].provenance == {
        "column_index": 1,
        "row_index": 1,
        "table_index": 1,
    }
    assert result.units[-1].text == "Tail"


def test_docx_output_is_deterministic() -> None:
    content = make_docx()
    descriptor = intake(
        filename="report.docx",
        content_type="application/zip",
        content=content,
    )
    extractor = DeterministicDocumentExtractor()

    first = extractor.extract(
        descriptor=descriptor,
        content=content,
    )
    second = extractor.extract(
        descriptor=descriptor,
        content=content,
    )

    assert first.to_public_dict() == second.to_public_dict()


def test_xlsx_preserves_sheet_and_cell_provenance() -> None:
    content = make_xlsx()
    descriptor = intake(
        filename="book.xlsx",
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        content=content,
    )

    result = DeterministicDocumentExtractor().extract(
        descriptor=descriptor,
        content=content,
    )

    assert result.parser == "openpyxl"
    assert result.unit_count == 5
    assert result.units[0].text == "Name"
    assert result.units[0].provenance[
        "coordinate"
    ] == "A1"
    assert result.units[3].text == "12.5"
    assert result.units[4].text == "=1+1"
    assert result.units[4].provenance[
        "sheet_name"
    ] == "Formula"


def test_xlsx_formula_is_not_recalculated() -> None:
    content = make_xlsx()
    descriptor = intake(
        filename="book.xlsx",
        content_type="application/zip",
        content=content,
    )

    result = DeterministicDocumentExtractor().extract(
        descriptor=descriptor,
        content=content,
    )

    assert any(
        unit.text == "=1+1"
        for unit in result.units
    )


def test_pdf_blank_page_has_page_provenance_warning() -> None:
    content = make_pdf()
    descriptor = intake(
        filename="blank.pdf",
        content_type="application/pdf",
        content=content,
    )

    result = DeterministicDocumentExtractor().extract(
        descriptor=descriptor,
        content=content,
    )

    assert result.unit_count == 1
    assert result.units[0].kind == (
        ExtractionUnitKind.PDF_PAGE
    )
    assert result.units[0].provenance == {
        "page_number": 1
    }
    assert result.units[0].text == ""
    assert result.warnings == (
        "PDF_PAGE_1_NO_TEXT",
    )


def test_unit_character_limit_fails_closed() -> None:
    content = b"123456"
    descriptor = intake(
        filename="notes.txt",
        content_type="text/plain",
        content=content,
    )
    extractor = DeterministicDocumentExtractor(
        DocumentExtractionPolicy(
            max_unit_characters=5,
        )
    )

    with pytest.raises(
        DocumentExtractionError,
    ) as captured:
        extractor.extract(
            descriptor=descriptor,
            content=content,
        )

    assert captured.value.code == (
        "DOCUMENT_EXTRACTION_UNIT_TOO_LARGE"
    )


def test_total_character_limit_fails_closed() -> None:
    content = make_docx()
    descriptor = intake(
        filename="report.docx",
        content_type="application/zip",
        content=content,
    )
    extractor = DeterministicDocumentExtractor(
        DocumentExtractionPolicy(
            max_total_characters=4,
        )
    )

    with pytest.raises(
        DocumentExtractionError,
    ) as captured:
        extractor.extract(
            descriptor=descriptor,
            content=content,
        )

    assert captured.value.code == (
        "DOCUMENT_EXTRACTION_CHARACTER_LIMIT"
    )


def test_unit_count_limit_fails_closed() -> None:
    content = make_xlsx()
    descriptor = intake(
        filename="book.xlsx",
        content_type="application/zip",
        content=content,
    )
    extractor = DeterministicDocumentExtractor(
        DocumentExtractionPolicy(
            max_units=2,
        )
    )

    with pytest.raises(
        DocumentExtractionError,
    ) as captured:
        extractor.extract(
            descriptor=descriptor,
            content=content,
        )

    assert captured.value.code == (
        "DOCUMENT_EXTRACTION_UNIT_LIMIT"
    )


def test_public_dict_can_exclude_text() -> None:
    content = b"hello"
    descriptor = intake(
        filename="notes.txt",
        content_type="text/plain",
        content=content,
    )

    result = DeterministicDocumentExtractor().extract(
        descriptor=descriptor,
        content=content,
    )
    public = result.to_public_dict(
        include_text=False
    )

    assert "text" not in public["units"][0]
    assert "text_sha256" in public["units"][0]


def test_policy_requires_positive_limits() -> None:
    with pytest.raises(ValueError):
        DocumentExtractionPolicy(
            max_units=0,
        )
